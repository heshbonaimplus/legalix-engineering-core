import docx, os, zipfile, re, shutil, json
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc_path = '/home/yogi/lod_project/מפרט_נתונים_לפיתוח_ספריית_חוקי_תכן_HVAC_Candidate_V1.docx'
doc = docx.Document()

def set_rtl(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

NAVY = RGBColor(16, 44, 87)
CRIMSON = RGBColor(180, 0, 0)
DARK_GRAY = RGBColor(50, 50, 50)
GREEN_COLOR = RGBColor(0, 120, 0)
ORANGE_COLOR = RGBColor(210, 105, 0)
BLUE_COLOR = RGBColor(0, 102, 204)

# Document Header Title
p_title = doc.add_paragraph()
set_rtl(p_title)
r_title = p_title.add_run('מפרט נתונים הנדסי לפיתוח — ספריית חוקי תכן MEP דיגיטלית (Candidate Rules V1)')
r_title.font.name = 'David'
r_title.font.size = Pt(18)
r_title.font.bold = True
r_title.font.color.rgb = NAVY

p_sub = doc.add_paragraph()
set_rtl(p_sub)
r_sub = p_sub.add_run('Legalix Engineering Architecture | מפרט טכני לצוות הפיתוח | סכמת 16 שדות | סטטוס: Candidate Rules V1 – Pending Engineering Verification')
r_sub.font.name = 'David'
r_sub.font.size = Pt(11.5)
r_sub.font.italic = True
r_sub.font.color.rgb = DARK_GRAY

# Section 1: Core Architecture & Exact System Statement
p_arch = doc.add_paragraph()
set_rtl(p_arch)
r_arch = p_arch.add_run('1. עקרון יסוד וארכיטקטורת 4 השכבות ההנדסיות (4-Layer Engine + Human Gate):')
r_arch.font.name = 'David'
r_arch.font.size = Pt(13.5)
r_arch.font.bold = True
r_arch.font.color.rgb = NAVY

p_arch_desc = doc.add_paragraph()
set_rtl(p_arch_desc)
arch_text = (
    "הצהרת יסוד הנדסית ומערכתית:\n"
    "\"ארכיטקטורה דטרמיניסטית לצמצום תלות בפרשנות AI ולמניעת שימוש ב־AI כמקור נורמטיבי.\"\n\n"
    "המערכת מבוססת על הפרדה חדה ומוחלטת בין 3 רכיבי הידע בכל בדיקה:\n"
    "1. מה ראיתי במודל (Geometry Engine): עובדות ומדידות בלבד מתוך RVT/IFC/DWG.\n"
    "2. מה חישבתי (Calculation Engine): נוסחאות פיזיקליות והידראוליות דטרמיניסטיות בלבד.\n"
    "3. מה התקן באמת דורש (Verified Rules Engine): השוואה מול ספי תקנים, טבלאות נתונים ומסמכים מאומתים בסעיף ומהדורה.\n"
    "4. שכבת AI Engineering Assistant: ניתוח משמעות החריגה והצעת חלופות הנדסיות ליועץ.\n"
    "5. שער אישור אנושי (Human Approval Gate): בקר התכן / היועץ המוסמך הוא הסמכות הבלעדית לאישור/דחיית ממצא או הכרזת Hold Point."
)
r_ad = p_arch_desc.add_run(arch_text)
r_ad.font.name = 'David'
r_ad.font.size = Pt(10.5)

# Status Matrix Table (4 Colors with refined GREEN and HOLD POINT policy)
p_stat = doc.add_paragraph()
set_rtl(p_stat)
r_stat = p_stat.add_run('2. מפתח 4 הסטטוסים ההנדסיים ומדיניות Recommended Hold Point:')
r_stat.font.name = 'David'
r_stat.font.size = Pt(13)
r_stat.font.bold = True
r_stat.font.color.rgb = NAVY

t_stat = doc.add_table(rows=5, cols=3)
t_stat.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['סטטוס וסיווג', 'הגדרה הנדסית מדויקת', 'הנחיית פעולה ליועץ / מודל']
hdr_row = t_stat.rows[0]
for idx, text in enumerate(headers):
    cell = hdr_row.cells[idx]
    shd = parse_xml('<w:shd {} w:fill="102C57"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shd)
    p = cell.paragraphs[0]
    set_rtl(p)
    r = p.add_run(text)
    r.font.name = 'David'
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)

status_data = [
    ('🔴 אדום (RED)', 'חובה לתקן — כשל קריטי, בטיחות חיים (Life Safety), התנגשות קשה בשלד או אי-עמידה בתקן מחייב.', 'חובת תיקון מיידי בתוכניות ובמודל ⛔'),
    ('🟡 צהוב (YELLOW)', 'דורש החלטת יועץ / תיאום — פערי תכן שניתן לגשר עליהם בפרקטיקת שטח, הנחיית תכנון או תיאום אדריכלי/קונסטרוקטיבי.', 'קבלת החלטת יועץ והגדרת מפרט ביצוע ⚠️'),
    ('🔵 כחול (BLUE)', 'המלצת אופטימיזציה / הנדסת ערך — המלצה לחיסכון בעלויות ליזם, יעילות אנרגטית ושדרוג מפרט.', 'המלצה כלכלית לבחירת היזם והמתכנן 💡'),
    ('🟢 ירוק (GREEN — PASS)', 'נבדק ונמצא תקין — לא נמצאה חריגה במסגרת הבדיקות האוטומטיות שהופעלו על האלמנט והנתונים הזמינים במודל.', 'נבדק ואומת במסגרת הבדיקות שהופעלו ✅')
]

for row_idx, (c1, c2, c3) in enumerate(status_data, start=1):
    row = t_stat.rows[row_idx]
    bg = 'FFF0F0' if '🔴' in c1 else ('FFFDF0' if '🟡' in c1 else ('F0F8FF' if '🔵' in c1 else 'F0FFF4'))
    tc = CRIMSON if '🔴' in c1 else (ORANGE_COLOR if '🟡' in c1 else (BLUE_COLOR if '🔵' in c1 else GREEN_COLOR))
    for col_idx, text in enumerate([c1, c2, c3]):
        cell = row.cells[col_idx]
        shd = parse_xml('<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg))
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        set_rtl(p)
        r = p.add_run(text)
        r.font.name = 'David'
        r.font.size = Pt(9.5)
        if col_idx == 0:
            r.font.bold = True
            r.font.color.rgb = tc

doc.add_paragraph()

# Policy Box for Recommended Hold Point
p_hp_pol = doc.add_paragraph()
set_rtl(p_hp_pol)
r_hpp = p_hp_pol.add_run('מדיניות המערכת לגבי נקודת עצירה (Recommended Hold Point Policy):\nהמערכת מפיקה התראת "RECOMMENDED HOLD POINT" (המלצה בלבד מטעם המערכת). חל איסור על התוכנה לקבוע עצירה משפטית אוטונומית — הכרזת Hold Point רשמית ומחייבת מתבצעת אך ורק בלחיצה ואישור של בקר התכן / היועץ המוסמך.')
r_hpp.font.name = 'David'
r_hpp.font.size = Pt(10)
r_hpp.font.italic = True
r_hpp.font.color.rgb = CRIMSON

doc.add_paragraph()

# Section 3: Refrigerant Lookup Table
p_ref = doc.add_paragraph()
set_rtl(p_ref)
r_ref = p_ref.add_run('3. טבלת נתוני גזי קירור נורמטיבית (Normative Refrigerant Dataset Table):')
r_ref.font.name = 'David'
r_ref.font.size = Pt(13)
r_ref.font.bold = True
r_ref.font.color.rgb = NAVY

t_ref = doc.add_table(rows=4, cols=4)
t_ref.alignment = WD_TABLE_ALIGNMENT.CENTER
ref_headers = ['סוג גז קירור (Refrigerant)', 'קבוצת בטיחות (Safety Group)', 'סף ריכוז מרבי (RCL kg/m³)', 'מקור תקן מחייב']
hdr_ref_row = t_ref.rows[0]
for idx, text in enumerate(ref_headers):
    cell = hdr_ref_row.cells[idx]
    shd = parse_xml('<w:shd {} w:fill="102C57"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shd)
    p = cell.paragraphs[0]
    set_rtl(p)
    r = p.add_run(text)
    r.font.name = 'David'
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)

ref_data = [
    ('R-410A (HFC Blend)', 'A1 (Non-Toxic / Non-Flammable)', '0.44 kg/m³', 'ת״י 920 / EN 378:2016 Table C.1 / ISO 5149'),
    ('R-32 (HFC Single Component)', 'A2L (Lower Flammability)', '0.06 kg/m³', 'ת״י 920 / EN 378:2016 Table C.1 / ISO 5149'),
    ('R-134a (HFC Single Component)', 'A1 (Non-Toxic / Non-Flammable)', '0.25 kg/m³', 'ת״י 920 / EN 378:2016 Table C.1 / ISO 5149')
]

for row_idx, (r1, r2, r3, r4) in enumerate(ref_data, start=1):
    row = t_ref.rows[row_idx]
    for col_idx, text in enumerate([r1, r2, r3, r4]):
        cell = row.cells[col_idx]
        shd = parse_xml('<w:shd {} w:fill="FFFFFF"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        set_rtl(p)
        r = p.add_run(text)
        r.font.name = 'David'
        r.font.size = Pt(9.5)
        if col_idx == 0:
            r.font.bold = True

doc.add_paragraph()

# Section 4: 16-Field Schema Table Master
p_rhead = doc.add_paragraph()
set_rtl(p_rhead)
r_rhead = p_rhead.add_run('4. ספרית חוקי התכן במבנה הסכמה המלא (Full 16-Field Data Specification):')
r_rhead.font.name = 'David'
r_rhead.font.size = Pt(13.5)
r_rhead.font.bold = True
r_rhead.font.color.rgb = NAVY

json_path = '/home/yogi/lod_project/HVAC_CANDIDATE_RULE_LIBRARY_V1.json'
with open(json_path, 'r', encoding='utf-8') as f:
    lib_data = json.load(f)

rules_list = lib_data['rules']

for r_idx, rule in enumerate(rules_list, start=1):
    rule_id = rule.get('Rule_ID', f'HVAC-RULE-{r_idx}')
    rule_name = rule.get('Rule_Name', 'כלל תכן')
    sev = rule.get('Severity', 'RED')
    
    p_rh = doc.add_paragraph()
    set_rtl(p_rh)
    r_rh = p_rh.add_run(f"מזהה כלל: [{rule_id}] — {rule_name}")
    r_rh.font.name = 'David'
    r_rh.font.size = Pt(12)
    r_rh.font.bold = True
    r_rh.font.color.rgb = CRIMSON if sev == 'RED' else (ORANGE_COLOR if sev == 'YELLOW' else BLUE_COLOR)
    
    t_rule = doc.add_table(rows=8, cols=2)
    t_rule.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    inputs_list = [f"{inp['name']} ({inp['type']})" for inp in rule.get('Required_Inputs', [])]
    inputs_str = ', '.join(inputs_list)
    
    schema_rows = [
        ("תחולת הכלל (Applicability) & אלמנט:", f"Applicability: {rule.get('Rule_Applicability')} | Element: {rule.get('Element_Type')}"),
        ("קלטים נדרשים (Required Inputs):", inputs_str),
        ("נוסחת חישוב & תנאי הדק (Logic & Trigger):", f"נוסחה: {rule.get('Calculation_Formula', 'Direct Geometry')}\nתנאי חריגה: {rule.get('Threshold_Condition')}"),
        ("סף נורמטיבי וסוג הסף (Threshold & Type):", f"סף: {rule.get('Standard_Requirement')}\nסוג סף: {rule.get('Threshold_Type')}"),
        ("מקור תקן, מהדורה וסעיף מדויק:", f"{rule.get('Standard_Source')} | מהדורה: {rule.get('Edition')} | סעיף: {rule.get('Clause')}"),
        ("סיווג נורמטיבי, חומרה ואישור אנושי:", f"סיווג: {rule.get('Requirement_Type')} | חומרה: {sev} | אישור אנושי נדרש: {rule.get('Human_Approval_Required')}"),
        ("רמות ביטחון (Confidence Levels):", f"זיהוי גיאומטרי (Detection Confidence): {rule.get('Detection_Confidence', '100%')} | אימות נורמטיבי (Source Confidence): {rule.get('Source_Confidence', '85%')} | סטטוס אימות: {rule.get('Rule_Verification')}"),
        ("סיווג פתרון והנחיית תיקון (Suggested Fix):", f"סוג פתרון: {rule.get('Suggested_Fix_Type')}\nהנחיה: {rule.get('Suggested_Remediation')}")
    ]
    
    for row_i, (k_txt, v_txt) in enumerate(schema_rows):
        row = t_rule.rows[row_i]
        
        # Left Label
        c0 = row.cells[0]
        shd0 = parse_xml('<w:shd {} w:fill="F4F6F9"/>'.format(nsdecls('w')))
        c0._tc.get_or_add_tcPr().append(shd0)
        p0 = c0.paragraphs[0]
        set_rtl(p0)
        r0 = p0.add_run(k_txt)
        r0.font.name = 'David'
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.color.rgb = NAVY
        c0.width = Inches(2.5)
        
        # Right Value
        c1 = row.cells[1]
        bg_c = 'FFF0F0' if (row_i == 5 and 'RED' in sev) else 'FFFFFF'
        shd1 = parse_xml('<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg_c))
        c1._tc.get_or_add_tcPr().append(shd1)
        p1 = c1.paragraphs[0]
        set_rtl(p1)
        r1 = p1.add_run(v_txt)
        r1.font.name = 'David'
        r1.font.size = Pt(9)
        if row_i == 5:
            r1.font.bold = True
            r1.font.color.rgb = CRIMSON if 'RED' in sev else ORANGE_COLOR
        c1.width = Inches(4.3)
        
    doc.add_paragraph() # Spacer

doc.save(doc_path)
print('Master Candidate V1 Data Specification DOCX generated successfully at:', doc_path)

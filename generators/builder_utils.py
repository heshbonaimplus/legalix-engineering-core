import docx, os, zipfile, re, shutil
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENTATION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_rtl(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.line_spacing = 1.25

NAVY = RGBColor(16, 44, 87)
CRIMSON = RGBColor(180, 0, 0)
DARK_GRAY = RGBColor(60, 60, 60)
GREEN_COLOR = RGBColor(0, 120, 0)
ORANGE_COLOR = RGBColor(210, 105, 0)
BLUE_COLOR = RGBColor(0, 102, 204)

def setup_clean_header_footer(s, title_str):
    # Ensure no duplicate headers in section
    hdr = s.header
    hdr.is_linked_to_previous = False
    for p in hdr.paragraphs:
        p.text = ""
    p_hdr = hdr.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hdr = p_hdr.add_run(title_str)
    r_hdr.font.name = 'David'
    r_hdr.font.size = Pt(8.5)
    r_hdr.font.color.rgb = RGBColor(120, 120, 120)
    
    ftr = s.footer
    ftr.is_linked_to_previous = False
    for p in ftr.paragraphs:
        p.text = ""
    p_ftr = ftr.paragraphs[0]
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_ftr = p_ftr.add_run("סבב בדיקה: 01 | מהדורה לקראת Rev 02 | Legalix Designer Correction & Closure")
    r_ftr.font.name = 'David'
    r_ftr.font.size = Pt(8.5)
    r_ftr.font.color.rgb = RGBColor(120, 120, 120)

print("Setup helper functions ready.")

import docx, os, zipfile, re, shutil, json
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc_path = '/home/yogi/lod_project/ספריית_חוקי_תכן_הנדסית_HVAC_מהדורת_V1_ייצור.docx'
doc = docx.Document()

def set_rtl(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

NAVY = RGBColor(16, 44, 87)
CRIMSON = RGBColor(180, 0, 0)
DARK_GRAY = RGBColor(50, 50, 50)
GREEN_COLOR = RGBColor(0, 120, 0)
ORANGE_COLOR = RGBColor(210, 105, 0)
BLUE_COLOR = RGBColor(0, 102, 204)

# Document Header Title
p_title = doc.add_paragraph()
set_rtl(p_title)
r_title = p_title.add_run('מפרט ספרית חוקי תכן מאומתת — Legalix Verified Rules Engine (HVAC V1)')
r_title.font.name = 'David'
r_title.font.size = Pt(18)
r_title.font.bold = True
r_title.font.color.rgb = NAVY

p_sub = doc.add_paragraph()
set_rtl(p_sub)
r_sub = p_sub.add_run('ארכיטקטורת 4 שכבות הנדסיות + שער אישור אנושי (Human Approval Gate) | מפרט טכני לפיתוח וייצור מסחרי')
r_sub.font.name = 'David'
r_sub.font.size = Pt(12)
r_sub.font.italic = True
r_sub.font.color.rgb = DARK_GRAY

# Section 1: 4 Architecture Layers + Human Gate
p_arch = doc.add_paragraph()
set_rtl(p_arch)
r_arch = p_arch.add_run('1. ארכיטקטורת 4 השכבות ההנדסיות ושער האישור האנושי (4-Layer Engine + Human Gate):')
r_arch.font.name = 'David'
r_arch.font.size = Pt(14)
r_arch.font.bold = True
r_arch.font.color.rgb = NAVY

p_arch_desc = doc.add_paragraph()
set_rtl(p_arch_desc)
arch_text = (
    "מערכת הבקרה מופרדת ל-4 שכבות עצמאיות המבטיחות אפס הזיות (Zero Hallucination) ומהימנות הנדסית משפטית:\n"
    "• שכבה 1 — מנוע גיאומטרי (Geometry Engine): חילוץ עובדות ומדידות גיאומטריות בלבד מתוך RVT/IFC/DWG (מרחקים, חתכים, גבהים, שיפועים).\n"
    "• שכבה 2 — מנוע חישוב הנדסי (Calculation Engine): ביצוע חישובים פיזיקליים והידראוליים דטרמיניסטיים (מהירות m/s, ספיקה m³/s, מפל לחץ Pa, ריכוז גז kg/m³, כוח פתיחת דלת N).\n"
    "• שכבה 3 — מנוע כללים מאומת (Verified Rules Engine): השוואה מול כללים מאומתים בלבד הכוללים סעיף תקן מדויק, שנת מהדורה וסיווג חובה/המלצה.\n"
    "• שכבה 4 — עוזר הנדסי AI (AI Engineering Assistant): מסביר את מהות הכשל, מחשב פערי חריגה, ומנסח הצעת תיקון כירורגית מותאמת.\n"
    "• שער אישור אנושי (Human Approval Gate): בקר התכן / היועץ המוסמך הוא הסמכות הבלעדית לאישור/דחיית ממצא או הכרזת Hold Point רשמית."
)
r_ad = p_arch_desc.add_run(arch_text)
r_ad.font.name = 'David'
r_ad.font.size = Pt(10.5)

# Status Matrix Table (4 Colors with refined GREEN and HOLD POINT policy)
p_stat = doc.add_paragraph()
set_rtl(p_stat)
r_stat = p_stat.add_run('2. מפתח 4 הסטטוסים ההנדסיים ומדיניות Hold Point:')
r_stat.font.name = 'David'
r_stat.font.size = Pt(13)
r_stat.font.bold = True
r_stat.font.color.rgb = NAVY

t_stat = doc.add_table(rows=5, cols=3)
t_stat.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['סטטוס וסיווג', 'הגדרה הנדסית מדויקת', 'הנחיית פעולה ליועץ / מודל']
hdr_row = t_stat.rows[0]
for idx, text in enumerate(headers):
    cell = hdr_row.cells[idx]
    shd = parse_xml('<w:shd {} w:fill="102C57"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shd)
    p = cell.paragraphs[0]
    set_rtl(p)
    r = p.add_run(text)
    r.font.name = 'David'
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)

status_data = [
    ('🔴 אדום (RED)', 'חובה לתקן — כשל קריטי, בטיחות חיים (Life Safety), התנגשות קשה בשלד או אי-עמידה בתקן מחייב.', 'חובת תיקון מיידי בתוכניות ובמודל ⛔'),
    ('🟡 צהוב (YELLOW)', 'דורש החלטת יועץ / תיאום — פערי תכן שניתן לגשר עליהם בפרקטיקת שטח, הנחיית תכנון או תיאום אדריכלי/קונסטרוקטיבי.', 'קבלת החלטת יועץ והגדרת מפרט ביצוע ⚠️'),
    ('🔵 כחול (BLUE)', 'המלצת אופטימיזציה / הנדסת ערך — המלצה לחיסכון בעלויות ליזם, יעילות אנרגטית ושדרוג מפרט.', 'המלצה כלכלית לבחירת היזם והמתכנן 💡'),
    ('🟢 ירוק (GREEN — PASS)', 'נבדק ונמצא תקין — לא נמצאה חריגה במסגרת הבדיקות האוטומטיות שהופעלו על האלמנט והנתונים הזמינים במודל.', 'נבדק ואומת במסגרת הבדיקות שהופעלו ✅')
]

for row_idx, (c1, c2, c3) in enumerate(status_data, start=1):
    row = t_stat.rows[row_idx]
    bg = 'FFF0F0' if '🔴' in c1 else ('FFFDF0' if '🟡' in c1 else ('F0F8FF' if '🔵' in c1 else 'F0FFF4'))
    tc = CRIMSON if '🔴' in c1 else (ORANGE_COLOR if '🟡' in c1 else (BLUE_COLOR if '🔵' in c1 else GREEN_COLOR))
    for col_idx, text in enumerate([c1, c2, c3]):
        cell = row.cells[col_idx]
        shd = parse_xml('<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg))
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        set_rtl(p)
        r = p.add_run(text)
        r.font.name = 'David'
        r.font.size = Pt(9.5)
        if col_idx == 0:
            r.font.bold = True
            r.font.color.rgb = tc

doc.add_paragraph()

# Policy Box for Recommended Hold Point
p_hp_pol = doc.add_paragraph()
set_rtl(p_hp_pol)
r_hpp = p_hp_pol.add_run('מדיניות המערכת לגבי נקודת עצירה (Recommended Hold Point Policy):\nהמערכת מפיקה התראת "RECOMMENDED HOLD POINT" (המלצה בלבד מטעם המערכת). חל איסור על התוכנה לקבוע עצירה משפטית אוטונומית — הכרזת Hold Point רשמית ומחייבת מתבצעת אך ורק בלחיצה ואישור של בקר התכן / היועץ המוסמך.')
r_hpp.font.name = 'David'
r_hpp.font.size = Pt(10)
r_hpp.font.italic = True
r_hpp.font.color.rgb = CRIMSON

doc.add_paragraph()

# Load Rules from JSON V1
json_path = '/home/yogi/lod_project/HVAC_RULE_LIBRARY_V1.json'
with open(json_path, 'r', encoding='utf-8') as f:
    lib_data = json.load(f)

# Rules Table Master (Full Production Schema)
p_rhead = doc.add_paragraph()
set_rtl(p_rhead)
r_rhead = p_rhead.add_run('3. ספרית חוקי התכן המאומתת לפיתוח וייצור (HVAC Verified Rules V1 Master):')
r_rhead.font.name = 'David'
r_rhead.font.size = Pt(14)
r_rhead.font.bold = True
r_rhead.font.color.rgb = NAVY

rules_list = lib_data['rules']
table_r = doc.add_table(rows=len(rules_list)+1, cols=7)
table_r.alignment = WD_TABLE_ALIGNMENT.CENTER
headers_r = ['Rule ID & סוג זיהוי', 'אלמנט ושם הכלל', 'קלט ונוסחת חישוב (Inputs & Formula)', 'סף ותנאי (Threshold & Condition)', 'מקור תקן, מהדורה וסעיף', 'סיווג תקן, חומרה ו-Confidence', 'הצעת תיקון (Suggested Fix)']
hdr_r_row = table_r.rows[0]
for idx, text in enumerate(headers_r):
    cell = hdr_r_row.cells[idx]
    shd = parse_xml('<w:shd {} w:fill="102C57"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shd)
    p = cell.paragraphs[0]
    set_rtl(p)
    r = p.add_run(text)
    r.font.name = 'David'
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)

for r_idx, rule in enumerate(rules_list, start=1):
    row = table_r.rows[r_idx]
    sev = rule.get('Severity', 'RED')
    bg = 'FFF0F0' if sev == 'RED' else ('FFFDF0' if sev == 'YELLOW' else ('F0F8FF' if sev == 'BLUE' else 'F0FFF4'))
    
    det_type = rule.get('Detection_Type', 'GEOMETRY')
    rule_id = rule.get('Rule_ID', f'HVAC-RULE-{r_idx}')
    rule_name = rule.get('Rule_Name', 'כלל תכן')
    elem_type = rule.get('Element_Type', 'Element')
    inputs_str = ', '.join(rule.get('Required_Inputs', []))
    formula_str = rule.get('Calculation_Formula', 'N/A')
    cond_str = rule.get('Trigger_Condition', 'N/A')
    thresh_str = rule.get('Standard_Threshold', 'N/A')
    src_str = rule.get('Standard_Source', 'ת״י')
    edit_str = rule.get('Edition', 'N/A')
    clause_str = rule.get('Clause', 'N/A')
    req_type = rule.get('Requirement_Type', 'Mandatory')
    conf_str = rule.get('Confidence', '100%')
    fix_str = rule.get('Suggested_Remediation', 'תיקון לפי תקן')
    
    col_vals = [
        f"{rule_id}\n[{det_type}]",
        f"{rule_name}\n({elem_type})",
        f"קלטים: {inputs_str}\nנוסחה: {formula_str}",
        f"תנאי: {cond_str}\nסף: {thresh_str}",
        f"תקן: {src_str}\nמהדורה: {edit_str}\nסעיף: {clause_str}",
        f"סיווג: {req_type}\nחומרה: {sev}\nביטחון: {conf_str}",
        fix_str
    ]
    
    for c_idx, val in enumerate(col_vals):
        cell = row.cells[c_idx]
        shd = parse_xml('<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg))
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        set_rtl(p)
        r = p.add_run(val)
        r.font.name = 'David'
        r.font.size = Pt(8.5)
        if c_idx == 0:
            r.font.bold = True
            r.font.color.rgb = NAVY
        elif c_idx == 5:
            r.font.bold = True
            r.font.color.rgb = CRIMSON if sev == 'RED' else (ORANGE_COLOR if sev == 'YELLOW' else BLUE_COLOR)

doc.save(doc_path)
print('Master Verified Rule Library V1 DOCX generated successfully at:', doc_path)

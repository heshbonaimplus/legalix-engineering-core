import docx, os, zipfile, re, shutil, json
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc_path = '/home/yogi/lod_project/ספריית_חוקי_תכן_הנדסית_מאסטר_אינסטלציה_וכיבוי.docx'
doc = docx.Document()

def set_rtl(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

NAVY = RGBColor(16, 44, 87)
CRIMSON = RGBColor(180, 0, 0)
DARK_GRAY = RGBColor(50, 50, 50)
GREEN_COLOR = RGBColor(0, 120, 0)
ORANGE_COLOR = RGBColor(210, 105, 0)
BLUE_COLOR = RGBColor(0, 102, 204)

# Document Title
p_title = doc.add_paragraph()
set_rtl(p_title)
r_title = p_title.add_run('ספריית חוקי תכן הנדסית מאסטר — אינסטלציה, ספרינקלרים וכיבוי אש (Plumbing & FP Rule Library)')
r_title.font.name = 'David'
r_title.font.size = Pt(18)
r_title.font.bold = True
r_title.font.color.rgb = NAVY

p_sub = doc.add_paragraph()
set_rtl(p_sub)
r_sub = p_sub.add_run('פרויקט: לוד ניר צבי (עמרם אברהם) | מנוע בקרת תכן MEP דיגיטלי | מודל 4 סטטוסים (אדום / צהוב / כחול / ירוק)')
r_sub.font.name = 'David'
r_sub.font.size = Pt(12)
r_sub.font.italic = True
r_sub.font.color.rgb = DARK_GRAY

# Overview of 3-Engine Architecture
p_arch = doc.add_paragraph()
set_rtl(p_arch)
r_arch = p_arch.add_run('ארכיטקטורת מנוע בקרת התכן הדיגיטלי (3-Engine Core Architecture):')
r_arch.font.name = 'David'
r_arch.font.size = Pt(14)
r_arch.font.bold = True
r_arch.font.color.rgb = NAVY

p_desc = doc.add_paragraph()
set_rtl(p_desc)
r_desc = p_desc.add_run('1. מנוע גיאומטריה (Geometry Engine): מודד בדיוק מילימטרי שיפועי קווים, קטרים, מרחקים מקורות שלד, שטחי כיסוי ספרינקלר, ומפלסי לולאות סניקה ישירות ממודלי RVT / IFC / DWG.\n2. מנוע כללים הנדסיים (Deterministic Rules Engine): משווה את המדידות מול ספי תכן ותקנים ישראליים (ת״י 1205, ת״י 1596, ת״י 1933, NFPA 13/14/20) במספרים מוחלטים (לדוגמה: IF Sprinkler.DistanceToBeam < Clearance -> RED).\n3. שכבת בינה מלאכותית (AI Synthesizer Layer): מנתחת את המשמעות ההידראולית והבטיחותית, מחשבת ספיקות ולחצים, ומייצרת הצעת תיקון ממוקדת ליועץ.\n4. מודל Digital Twin & BCF: הצגת הממצאים ישירות על גבי המודל התלת-ממדי ברוויט עם אפשרות הכרעה ליועץ (אישור / דחייה / אימוץ המלצה).')
r_desc.font.name = 'David'
r_desc.font.size = Pt(10.5)

# Status Model (4 Colors)
p_stat = doc.add_paragraph()
set_rtl(p_stat)
r_stat = p_stat.add_run('מפתח 4 הסטטוסים ההנדסיים (4-Tier Status Model):')
r_stat.font.name = 'David'
r_stat.font.size = Pt(13)
r_stat.font.bold = True
r_stat.font.color.rgb = NAVY

t_stat = doc.add_table(rows=5, cols=3)
t_stat.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['סטטוס וסיווג', 'הגדרה הנדסית ומשמעות', 'הנחיית פעולה ליועץ / מודל']
hdr_row = t_stat.rows[0]
for idx, text in enumerate(headers):
    cell = hdr_row.cells[idx]
    shd = parse_xml('<w:shd {} w:fill="102C57"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shd)
    p = cell.paragraphs[0]
    set_rtl(p)
    r = p.add_run(text)
    r.font.name = 'David'
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)

status_data = [
    ('🔴 אדום (RED)', 'חובה לתקן — כשל קריטי, בטיחות חיים, התנגשות קשה בשלד, סכנת הצפת ביוב או פסילת כב״ה/פקע״ר.', 'חובת תיקון מיידי בתוכניות ובמודל ⛔'),
    ('🟡 צהוב (YELLOW)', 'דורש החלטת יועץ / תיאום ביצוע מול קונסטרוקציה, מפרט בטון אטום במאגרים או חלופת שטח מקובלת.', 'קבלת החלטת יועץ והגדרת מפרט ביצוע ⚠️'),
    ('🔵 כחול (BLUE)', 'המלצת אופטימיזציה, הנדסת ערך (Value Engineering), שדרוג לראשי K=8.0 וחיסכון בעלויות משאבות ליזם.', 'המלצה כלכלית לבחירת היזם והמתכנן 💡'),
    ('🟢 ירוק (GREEN)', 'נבדק במודל ואומת כתקין ב-100% לפי כל דרישות התקנים והחישובים ההידראוליים.', 'נבדק ואושר לתעודת גמר / טופס 4 ✅')
]

for row_idx, (c1, c2, c3) in enumerate(status_data, start=1):
    row = t_stat.rows[row_idx]
    bg = 'FFF0F0' if '🔴' in c1 else ('FFFDF0' if '🟡' in c1 else ('F0F8FF' if '🔵' in c1 else 'F0FFF4'))
    tc = CRIMSON if '🔴' in c1 else (ORANGE_COLOR if '🟡' in c1 else (BLUE_COLOR if '🔵' in c1 else GREEN_COLOR))
    for col_idx, text in enumerate([c1, c2, c3]):
        cell = row.cells[col_idx]
        shd = parse_xml('<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg))
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        set_rtl(p)
        r = p.add_run(text)
        r.font.name = 'David'
        r.font.size = Pt(9.5)
        if col_idx == 0:
            r.font.bold = True
            r.font.color.rgb = tc

doc.add_paragraph()

# Load Rules from JSON
json_path = '/home/yogi/lod_project/PLUMBING_RULE_LIBRARY.json'
with open(json_path, 'r', encoding='utf-8') as f:
    lib_data = json.load(f)

# Rules Table
p_rhead = doc.add_paragraph()
set_rtl(p_rhead)
r_rhead = p_rhead.add_run('ספריית חוקי התכן המובנית לאינסטלציה וכיבוי אש (Plumbing & FP Rules Schema):')
r_rhead.font.name = 'David'
r_rhead.font.size = Pt(14)
r_rhead.font.bold = True
r_rhead.font.color.rgb = NAVY

rules_list = lib_data['rules']
table_r = doc.add_table(rows=len(rules_list)+1, cols=6)
table_r.alignment = WD_TABLE_ALIGNMENT.CENTER
headers_r = ['מזהה חוק (Rule ID)', 'שם הכלל ותיאור הכשל', 'תנאי בדיקה גיאומטרי/הנדסי (Condition & Threshold)', 'חומרה (Severity)', 'מקור תקן (Standard Source)', 'הצעת תיקון (Remediation)']
hdr_r_row = table_r.rows[0]
for idx, text in enumerate(headers_r):
    cell = hdr_r_row.cells[idx]
    shd = parse_xml('<w:shd {} w:fill="102C57"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shd)
    p = cell.paragraphs[0]
    set_rtl(p)
    r = p.add_run(text)
    r.font.name = 'David'
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)

for r_idx, rule in enumerate(rules_list, start=1):
    row = table_r.rows[r_idx]
    sev = rule['Severity']
    bg = 'FFF0F0' if sev == 'RED' else ('FFFDF0' if sev == 'YELLOW' else ('F0F8FF' if sev == 'BLUE' else 'F0FFF4'))
    
    col_vals = [
        rule['Rule_ID'],
        f"{rule['Rule_Name']}\n({rule['Element_Type']})",
        f"תנאי: {rule['Trigger_Condition']}\nסף: {rule['Standard_Threshold']}",
        f"🔴 אדום\nחובה לתקן" if sev == 'RED' else (f"🟡 צהוב\nתיאום יועץ" if sev == 'YELLOW' else f"🔵 כחול\nאופטימיזציה"),
        rule['Standard_Source'],
        rule['Suggested_Remediation']
    ]
    
    for c_idx, val in enumerate(col_vals):
        cell = row.cells[c_idx]
        shd = parse_xml('<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg))
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        set_rtl(p)
        r = p.add_run(val)
        r.font.name = 'David'
        r.font.size = Pt(8.5)
        if c_idx == 0:
            r.font.bold = True
            r.font.color.rgb = NAVY
        elif c_idx == 3:
            r.font.bold = True
            r.font.color.rgb = CRIMSON if sev == 'RED' else (ORANGE_COLOR if sev == 'YELLOW' else BLUE_COLOR)

doc.save(doc_path)
print('Master Plumbing Rule Library DOCX generated successfully at:', doc_path)

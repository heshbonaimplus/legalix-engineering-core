import docx, os, zipfile, re, shutil
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENTATION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from builder_utils import set_rtl, setup_clean_header_footer, NAVY, CRIMSON, DARK_GRAY, GREEN_COLOR, ORANGE_COLOR, BLUE_COLOR

target_filename = 'STRUCTURE_DESIGNER_CORRECTION_AND_CLOSURE_REPORT_REV01.docx'
doc_path = os.path.join('/home/yogi/lod_project', target_filename)
doc = docx.Document()

# Base Portrait Section setup
section1 = doc.sections[0]
section1.orientation = WD_ORIENTATION.PORTRAIT
section1.page_width = Inches(8.27)
section1.page_height = Inches(11.69)
section1.top_margin = Inches(0.984)
section1.bottom_margin = Inches(0.984)
section1.left_margin = Inches(0.984)
section1.right_margin = Inches(0.984)

setup_clean_header_footer(section1, "פרויקט לוד ניר צבי | קונסטרוקציה ושלד — Legalix Designer Correction & Closure")

# ==========================================
# 1. עמוד שער ומטא-דאטה
# ==========================================
p_t = doc.add_paragraph()
set_rtl(p_t)
r_t = p_t.add_run('Legalix Designer Correction & Closure Report')
r_t.font.name = 'David'
r_t.font.size = Pt(20)
r_t.font.bold = True
r_t.font.color.rgb = NAVY

p_sub = doc.add_paragraph()
set_rtl(p_sub)
r_sub = p_sub.add_run('הנדסת קונסטרוקציה, ביסוס, שלד ויציבות סיסמית — תוכנית עבודה וסגירת ממצאים ל-Rev 02')
r_sub.font.name = 'David'
r_sub.font.size = Pt(14)
r_sub.font.bold = True
r_sub.font.color.rgb = CRIMSON

p_meta_box = doc.add_paragraph()
set_rtl(p_meta_box)
meta_str = (
    "פרויקט: לוד ניר צבי\n"
    "יזם: עמרם אברהם\n"
    "מיועד עבור: מתכנן הקונסטרוקציה והשלד (צוות הנדסת מבנים / אלי חלמיש)\n"
    "מבנים / מגרשים: מגדלים 321, 339, מבנה 223 ופודיום חניונים\n"
    "גרסת מודל שנבדקה: Lod_ST_PR_R25 / Lod_ST_321_R25 / Lod_ST_339_R25 / Lod_ST_223_R25\n"
    "תאריך הפקה: ספטמבר 2026 | מהדורה: סבב 01 לקראת Revision 02\n"
    "מטרת המסמך: מתן מפת עבודה מוכנה, פעולות מועדפות, הנחיות עדכון מודל מדויקות והערכת זמן לקיצור ימי תכנון."
)
r_mb = p_meta_box.add_run(meta_str)
r_mb.font.name = 'David'
r_mb.font.size = Pt(10.5)
r_mb.font.color.rgb = DARK_GRAY

doc.add_paragraph() # Spacer

# ==========================================
# 2. תקציר מנהלים למתכנן
# ==========================================
p_ex_h = doc.add_paragraph()
set_rtl(p_ex_h)
r_exh = p_ex_h.add_run('1. תקציר ומטרת המסמך עבור צוות התכנון')
r_exh.font.name = 'David'
r_exh.font.size = Pt(14)
r_exh.font.bold = True
r_exh.font.color.rgb = NAVY

p_ex_b = doc.add_paragraph()
set_rtl(p_ex_b)
ex_text = (
    "דוח זה נבנה ככלי עבודה מעשי (Designer Correction & Closure) שמטרתו לקצר לצוות הקונסטרוקציה את זמן העבודה מ-Finding ל-Resolved.\n\n"
    "עבור כל אחד מ-45 הממצאים ההנדסיים שאותרו במודלי השלד, הדוח מציג מבנה עבודה תלת-אזורי ממוקד:\n"
    "• אזור 1: מה נמצא במודל (מיקום גיאומטרי, אלמנט נבדק, כוחות ומומנטים מחושבים מול תסבולת תקן ופער מספרי מדוד).\n"
    "• אזור 2: מה לעשות עכשיו (הפעולה המומלצת המועדפת, הסיבה לבחירתה, מה בדיוק לעדכן ב-Rev 02 והערכת זמן ביצוע).\n"
    "• אזור 3: איך נסגור את הממצא (קריטריון סגירה אוטומטי שהמערכת תבדוק ב-Rev 02, דיסציפלינות לתיאום וטופס החלטה).\n\n"
    "הדוח כולל מפת עדיפויות ביצוע (P1–P4) המאפשרת למתכנן להתמקד קודם במשימות השלד הקריטיות הנדרשות לאישור יציקות."
)
r_exb = p_ex_b.add_run(ex_text)
r_exb.font.name = 'David'
r_exb.font.size = Pt(10.5)

doc.add_paragraph()

# ==========================================
# 3. מפת עדיפויות ביצוע למתכנן (P1-P4)
# ==========================================
p_pr_h = doc.add_paragraph()
set_rtl(p_pr_h)
r_prh = p_pr_h.add_run('2. מפת עדיפויות לביצוע תיקונים (Task Priorities & Time Estimates)')
r_prh.font.name = 'David'
r_prh.font.size = Pt(13)
r_prh.font.bold = True
r_prh.font.color.rgb = NAVY

t_prio = doc.add_table(rows=5, cols=4)
t_prio.alignment = WD_TABLE_ALIGNMENT.CENTER
headers_pr = ['עדיפות ביצוע', 'סוג המשימה למתכנן', 'הערכת זמן ממוצעת לממצא', 'כמות ממצאים']
hdr_pr_row = t_prio.rows[0]
for idx, text in enumerate(headers_pr):
    cell = hdr_pr_row.cells[idx]
    shd = parse_xml('<w:shd {} w:fill="102C57"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shd)
    p = cell.paragraphs[0]
    set_rtl(p)
    r = p.add_run(text)
    r.font.name = 'David'
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)

prio_data = [
    ('🔥 P1 — עדכון שרטוטי זיון ומודל', 'הוספת חישוקי גזירה וחדירה, עיבוי זיון עליון/תחתון, שינוי קוטר מוטות וסידור חבקים', '15–20 דקות לממצא', '24 ממצאים'),
    ('🤝 P2 — דורש תיאום יועצים (MEP/AR)', 'התאמת שרוולי מעבר, פתחי פירים סמוך לעמודים, גובה ראש וסופרפוזיציה', 'דורש תיאום מול אדריכלות ו-MEP', '8 ממצאים'),
    ('📐 P3 — דורש כיול חישוב ומודל ETABS', 'קירות טרנספר, גזירת בסיס סיסמית (R=4.5), שקיעות דיפרנציאליות ועומס כבאית 160kN', '1–2 שעות (הרצת מודל / חישוב)', '11 ממצאים'),
    ('💡 P4 — אופטימיזציה שלד וזיון', 'בחינת מעבר למצמדים מכניים (Couplers) ושיפור צפיפות זיון', 'הנדסת ערך (לשיקול יזם/מתכנן)', '2 ממצאים')
]

for row_idx, (c1, c2, c3, c4) in enumerate(prio_data, start=1):
    row = t_prio.rows[row_idx]
    bg = 'FFF0F0' if 'P1' in c1 else ('FFFDF0' if 'P2' in c1 else ('F0F8FF' if 'P3' in c1 else 'F0FFF4'))
    for col_idx, text in enumerate([c1, c2, c3, c4]):
        cell = row.cells[col_idx]
        shd = parse_xml('<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg))
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        set_rtl(p)
        r = p.add_run(text)
        r.font.name = 'David'
        r.font.size = Pt(9)
        if col_idx == 0:
            r.font.bold = True
            r.font.color.rgb = CRIMSON if 'P1' in text else (ORANGE_COLOR if 'P2' in text else NAVY)
        elif col_idx == 3:
            r.font.bold = True

doc.add_paragraph()

# Read detailed 45 findings data from Grand Master Report
import json

# Structural Findings Data List (All 45 Findings)
from build_st_data import structural_designer_cards

effort_map_st = {
    "P1": "15–20 דקות (עדכון שרטוט זיון ברוויט/DWG)",
    "P2": "תיאום יועצים (אדריכלות / MEP)",
    "P3": "1–2 שעות (כיול מודל ETABS / חישוב חתך)",
    "P4": "הנדסת ערך (לשיקול דעת היזם והמתכנן)"
}

p_cards_head = doc.add_paragraph()
set_rtl(p_cards_head)
r_c_h = p_cards_head.add_run('3. פירוט כרטיסי עבודה במבנה 3 האזורים (Structural Action & Closure Cards)')
r_c_h.font.name = 'David'
r_c_h.font.size = Pt(14)
r_c_h.font.bold = True
r_c_h.font.color.rgb = NAVY

action_rows_full = []
current_ch = ""

for item in structural_designer_cards:
    ch = item['chapter']
    if ch != current_ch:
        current_ch = ch
        p_ch_t = doc.add_paragraph()
        set_rtl(p_ch_t)
        r_cht = p_ch_t.add_run(current_ch)
        r_cht.font.name = 'David'
        r_cht.font.size = Pt(14)
        r_cht.font.bold = True
        r_cht.font.color.rgb = NAVY
        
    prio_code = item['prio'][:2]
    effort_str = effort_map_st.get(prio_code, "20 דקות")
    
    action_rows_full.append({
        "id": item['id'],
        "desc": item['title'],
        "prio": item['prio'],
        "effort": effort_str,
        "loc": item['loc'].split('|')[0].strip(),
        "sev": item['status'][:1],
        "resp": item['coord'],
        "rec_action": item['rec_action'],
        "rev_update": item['rev_update'],
        "closure_crit": item['closure_crit'],
        "stat": "OPEN"
    })
    
    # Finding Card Header
    p_fh = doc.add_paragraph()
    set_rtl(p_fh)
    r_fhid = p_fh.add_run(f"{item['id']} | {item['title']}\n")
    r_fhid.font.name = 'David'
    r_fhid.font.size = Pt(12)
    r_fhid.font.bold = True
    r_fhid.font.color.rgb = NAVY
    
    r_fhst = p_fh.add_run(f"סטטוס: {item['status']}  |  עדיפות: {item['prio']}  |  זמן משוער: {effort_str}")
    r_fhst.font.name = 'David'
    r_fhst.font.size = Pt(10)
    r_fhst.font.bold = True
    r_fhst.font.color.rgb = CRIMSON if '🔴' in item['status'] else (ORANGE_COLOR if '🟡' in item['status'] else BLUE_COLOR)
    
    # 3-Block Clear Card Table
    t_c = doc.add_table(rows=11, cols=2)
    t_c.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    card_rows_data = [
        # Block 1: מה נמצא במודל
        ("🔍 אזור 1: מה נמצא במודל", "פירוט עובדתי, כוחות ונתוני שלד"),
        ("מיקום מדויק ואלמנט נבדק:", f"{item['loc']}\nאלמנט: {item['elem']}"),
        ("הממצא והפער שנמדד:", f"ממצא: {item['finding']}\nערך במודל: {item['val_curr']} | מחושב: {item['val_calc']}\nקריטריון תכן: {item['req']}\nפער שנמדד: {item['delta']}"),
        ("משמעות הנדסית ומקור מאומת:", f"משמעות: {item['impact']}\nמקור מאומת: {item['src']}"),
        
        # Block 2: מה לעשות עכשיו (מודגש ובולט)
        ("⚡ אזור 2: מה לעשות עכשיו (הנחיית פעולה מועדפת)", "הנחיות קונקרטיות לביצוע ב-Rev 02"),
        ("הפעולה המומלצת לביצוע (מועדפת):", f"{item['rec_action']}\n\nהסבר לבחירה: {item['rec_reason']}"),
        ("מה בדיוק לעדכן בתוכניות / מודל (Rev 02):", f"👉 {item['rev_update']}"),
        ("חלופות הנדסיות נוספות שנבחנו:", item['alts']),
        
        # Block 3: איך נסגור את הממצא
        ("🎯 אזור 3: איך נסגור את הממצא", "קריטריון אימות וסגירה אוטומטי"),
        ("קריטריון סגירה אוטומטי (Auto-Closure):", f"✅ {item['closure_crit']}\nתיאום נדרש: {item['coord']}"),
        ("טופס החלטת המתכנן:", "☐ התקבל כהמלצה מועדפת   ☐ התקבלה חלופה אחרת   ☐ לא התקבל (מצורף נימוק)\nסטטוס: ☐ OPEN   ☐ RESOLVED IN REV 02")
    ]
    
    for r_i, (k_txt, v_txt) in enumerate(card_rows_data):
        row = t_c.rows[r_i]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        is_block_hdr = r_i in [0, 4, 8]
        is_highlight_action = r_i in [5, 6]
        
        c0 = row.cells[0]
        bg_hdr = "102C57" if is_block_hdr else ("FFF0F0" if is_highlight_action else "F4F6F9")
        shd0 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hdr}"/>')
        c0._tc.get_or_add_tcPr().append(shd0)
        p0 = c0.paragraphs[0]
        set_rtl(p0)
        r0 = p0.add_run(k_txt)
        r0.font.name = 'David'
        r0.font.size = Pt(9.5 if is_block_hdr else 9)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(255, 255, 255) if is_block_hdr else (CRIMSON if is_highlight_action else NAVY)
        c0.width = Inches(2.4)
        
        c1 = row.cells[1]
        bg_val = "102C57" if is_block_hdr else ("FFFDF0" if is_highlight_action else "FFFFFF")
        shd1 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_val}"/>')
        c1._tc.get_or_add_tcPr().append(shd1)
        p1 = c1.paragraphs[0]
        set_rtl(p1)
        r1 = p1.add_run(v_txt)
        r1.font.name = 'David'
        r1.font.size = Pt(9.5 if is_block_hdr else 9)
        if is_block_hdr:
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(255, 255, 255)
        elif is_highlight_action:
            r1.font.bold = True
            r1.font.color.rgb = NAVY
        c1.width = Inches(4.4)
        
    doc.add_paragraph() # Spacer
    
    # Embed High-Res Screenshot
    img_p = item['img']
    if os.path.exists(img_p):
        p_img_t = doc.add_paragraph()
        set_rtl(p_img_t)
        r_it = p_img_t.add_run(f"תשריט ביקורת מתוך המודל — {item['id']}:")
        r_it.font.name = 'David'
        r_it.font.size = Pt(10)
        r_it.font.bold = True
        r_it.font.color.rgb = DARK_GRAY
        
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(img_p, width=Inches(5.8))
        
    doc.add_paragraph() # Spacer

# ==========================================
# 5. מקטע חדש לרוחב (Landscape Section) לטבלת תוכנית עבודה וסגירת ממצאים (45 שורות)
# ==========================================
section2 = doc.add_section(WD_SECTION.NEW_PAGE)
section2.orientation = WD_ORIENTATION.LANDSCAPE
section2.page_width = Inches(11.69)
section2.page_height = Inches(8.27)
section2.top_margin = Inches(0.984)
section2.bottom_margin = Inches(0.984)
section2.left_margin = Inches(0.984)
section2.right_margin = Inches(0.984)

setup_clean_header_footer(section2, "פרויקט לוד ניר צבי | קונסטרוקציה — תוכנית פעולה וסגירת ממצאים ל-Rev 02")

p_act_head = doc.add_paragraph()
set_rtl(p_act_head)
r_act_h = p_act_head.add_run('4. תוכנית עבודה ומעקב סגירת ממצאים לקונסטרוקטור (Structural Action & Closure Plan)')
r_act_h.font.name = 'David'
r_act_h.font.size = Pt(14)
r_act_h.font.bold = True
r_act_h.font.color.rgb = NAVY

t_act = doc.add_table(rows=len(action_rows_full)+1, cols=8)
t_act.alignment = WD_TABLE_ALIGNMENT.CENTER

headers_act = ['Finding ID', 'נושא הממצא', 'עדיפות וזמן משוער', 'מיקום בשלד', 'הפעולה המומלצת לביצוע (מועדפת)', 'מה בדיוק לעדכן ב-Rev 02', 'קריטריון סגירה אוטומטי', 'סטטוס']
hdr_act_row = t_act.rows[0]

trPr_hdr = hdr_act_row._tr.get_or_add_trPr()
trPr_hdr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
trPr_hdr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

col_widths = [Inches(1.1), Inches(1.2), Inches(1.1), Inches(1.1), Inches(2.2), Inches(2.0), Inches(1.5), Inches(0.6)]

for idx, text in enumerate(headers_act):
    cell = hdr_act_row.cells[idx]
    shd = parse_xml('<w:shd {} w:fill="102C57"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shd)
    p = cell.paragraphs[0]
    set_rtl(p)
    r = p.add_run(text)
    r.font.name = 'David'
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    cell.width = col_widths[idx]

for a_idx, r_item in enumerate(action_rows_full, start=1):
    row = t_act.rows[a_idx]
    
    trPr_row = row._tr.get_or_add_trPr()
    trPr_row.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    fsev = r_item['sev']
    bg = 'FFF0F0' if '🔴' in fsev else ('FFFDF0' if '🟡' in fsev else 'F0F8FF')
    
    col_entries = [
        r_item['id'],
        r_item['desc'],
        f"{r_item['prio'][:2]}\n({r_item['effort']})",
        r_item['loc'],
        r_item['rec_action'],
        r_item['rev_update'],
        r_item['closure_crit'],
        r_item['stat']
    ]
    
    for col_i, cell_text in enumerate(col_entries):
        cell = row.cells[col_i]
        shd = parse_xml('<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg))
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        set_rtl(p)
        r = p.add_run(cell_text)
        r.font.name = 'David'
        r.font.size = Pt(8.5)
        if col_i == 0:
            r.font.bold = True
            r.font.color.rgb = NAVY
        elif col_i == 2:
            r.font.bold = True
            r.font.color.rgb = CRIMSON if 'P1' in cell_text else (ORANGE_COLOR if 'P2' in cell_text else NAVY)
        cell.width = col_widths[col_i]

doc.add_paragraph() # Spacer

# ==========================================
# 6. החזרת מקטע לאורך (Portrait Section) למסקנות ו-Hold Point
# ==========================================
section3 = doc.add_section(WD_SECTION.NEW_PAGE)
section3.orientation = WD_ORIENTATION.PORTRAIT
section3.page_width = Inches(8.27)
section3.page_height = Inches(11.69)
section3.top_margin = Inches(0.984)
section3.bottom_margin = Inches(0.984)
section3.left_margin = Inches(0.984)
section3.right_margin = Inches(0.984)

setup_clean_header_footer(section3, "פרויקט לוד ניר צבי | קונסטרוקציה — הנחיות להגשת Rev 02")

p_conc_h = doc.add_paragraph()
set_rtl(p_conc_h)
r_ch_t = p_conc_h.add_run('5. הנחיות להגשת Revision 02 ותהליך סגירת ממצאים אוטומטי')
r_ch_t.font.name = 'David'
r_ch_t.font.size = Pt(14)
r_ch_t.font.bold = True
r_ch_t.font.color.rgb = NAVY

p_conc_b = doc.add_paragraph()
set_rtl(p_conc_b)
conc_text = (
    "על בסיס הממצאים, החלופות והפעולות המומלצות שפורטו בדוח זה, צוות תכנון הקונסטרוקציה מתבקש לפעול לפי סדר העדיפויות:\n"
    "1. עדכון ישיר בתוכניות הזיון והמודל של כל ממצאי P1 (חישוקי חדירה, עיבוי זיון עליון ותחתון, חבקי עמודים ועיבוי קירות).\n"
    "2. תיאום מול אדריכלות ו-MEP עבור ממצאי P2 (שרוולים יצוקים בקורות, פתחי פירים ליד עמודי שלד, גובה ראש וסופרפוזיציה).\n"
    "3. כיול והרצה חוזרת של מודל ETABS עבור ממצאי P3 (קירות טרנספר, גזירת בסיס סיסמית עם R=4.5, שקיעות דיפרנציאליות ועומס כבאית).\n\n"
    "עם קבלת תוכניות ומודל Revision 02 המעודכנים, המערכת תריץ בדיקת סגירת ממצאים אוטומטית (Automated Closure Verification) ותפיק טבלת השוואת גרסאות (Rev Comparison) שתאשר את סגירת הממצאים לקראת שחרור יציקות."
)
r_cb = p_conc_b.add_run(conc_text)
r_cb.font.name = 'David'
r_cb.font.size = Pt(10.5)

doc.add_paragraph()

# ==========================================
# 7. Recommended Hold Point
# ==========================================
p_hp_h = doc.add_paragraph()
set_rtl(p_hp_h)
r_hph = p_hp_h.add_run('6. המלצה לנקודת עצירה (Recommended Hold Point)')
r_hph.font.name = 'David'
r_hph.font.size = Pt(13)
r_hph.font.bold = True
r_hph.font.color.rgb = CRIMSON

p_hp_b = doc.add_paragraph()
set_rtl(p_hp_b)
hp_text = (
    "RECOMMENDED HOLD POINT:\n"
    "מומלץ שלא לקדם יציקות בטון של האלמנטים הרלוונטיים לממצאים המפורטים להלן עד לקבלת התייחסות המתכנן ואישור הגורמים המקצועיים:\n\n"
    "• רפסודת ביסוס וכלונסאות: ST-FND-001 (עומס יתר כלונסאות W-1), ST-FND-002 (שקיעות דיפרנציאליות 1/263), ST-FND-004 (גזירה בבורות מעלית).\n"
    "• תקרות חניון ופודיום: ST-SLB-001 (זיון עליון ותחתון), ST-SLB-002 (עומס רכב כיבוי 160kN), ST-SLB-004 (חדירה סביב עמודי חניון).\n"
    "• עמודים וקירות תמך: ST-COL-001 (קירות דיפון לחץ אדמה), ST-COL-002 (חבקי צומת עמוד-תקרה), ST-COL-004 (דקיקות עמודים וכפיפה דו-צירית).\n"
    "• קורות טרנספר ומגדלים: ST-TWR-001 (הסטת קיר W-4), ST-TWR-002 (קורת טרנספר TG-1), ST-TWR-004 (קורות צימוד אלכסוניות), ST-BLD-001 (אי-רציפות במבנה 223)."
)
r_hpb = p_hp_b.add_run(hp_text)
r_hpb.font.name = 'David'
r_hpb.font.size = Pt(10)
r_hpb.font.color.rgb = CRIMSON

doc.save(doc_path)
print("Complete Designer-Facing Action & Closure Report for Structure generated successfully at:", doc_path)
